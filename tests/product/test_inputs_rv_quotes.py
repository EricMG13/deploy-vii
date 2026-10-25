"""Task RV-1: the user-supplied relative-value quote sheet (`rv_quotes`) and its validator.

Workbooks here are written by openpyxl, an independent OOXML writer, so the validator is
exercised on bytes it did not produce. Every row is synthetic: a fictional issuer and
fictional providers, no market data.
"""

import io
import shutil
import sys
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Any
from xml.sax.saxutils import quoteattr
from zipfile import ZIP_DEFLATED, ZIP_STORED, ZipFile

import openpyxl
import pytest
from hypothesis import given, settings
from hypothesis import strategies as st

from deploy_vii.contracts import Run, Span
from deploy_vii.inputs.rv_quotes import (
    COLUMNS,
    FIX,
    LETTERS,
    LEVELS,
    MEDIA_TYPE,
    Issue,
    ValidationReport,
    sheet_cells,
    validate_rv_bytes,
    validate_rv_quotes,
)
from deploy_vii.preparation import PreparationRequest, SourceRequest, load_pack, prepare
from scripts.build_templates import README_LINES, SYNTHETIC_ROWS, build, workbook_bytes

ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "docs" / "templates" / "rv-quotes.template.xlsx"

FILE = "quotes.xlsx"
# The brief's contract, stated here independently of the module's constants.
BRIEF_COLUMNS = (
    "provider",
    "as_of",
    "issuer",
    "facility_id",
    "instrument",
    "quote_kind",
    "price",
    "yield_pct",
    "spread_bps",
    "bid",
    "ask",
    "currency",
    "note",
)
REQUIRED_COLUMNS = (
    "provider",
    "as_of",
    "issuer",
    "facility_id",
    "instrument",
    "quote_kind",
    "currency",
)
HEADER_FIX = (
    "set row 1 to exactly: provider, as_of, issuer, facility_id, instrument, quote_kind, "
    "price, yield_pct, spread_bps, bid, ask, currency, note (A1 to M1, nothing after M1)"
)
VALID: dict[str, Any] = {
    "provider": "Synthetic Dealer A",
    "as_of": "2026-09-30",
    "issuer": "Example Holdings Ltd (synthetic)",
    "facility_id": "TLB-2030",
    "instrument": "Term Loan B due 2030",
    "quote_kind": "observed",
    "price": 98.75,
    "yield_pct": None,
    "spread_bps": None,
    "bid": None,
    "ask": None,
    "currency": "USD",
    "note": "synthetic example row",
}


def write(
    directory: Path,
    rows: list[dict[str, Any]],
    *,
    header: tuple[str, ...] = COLUMNS,
    sheet: str = "rv_quotes",
) -> Path:
    book = openpyxl.Workbook()
    worksheet = book.active
    assert worksheet is not None
    worksheet.title = sheet
    if header:
        worksheet.append(list(header))
    for row in rows:
        worksheet.append([row.get(column) for column in header or COLUMNS])
    path = directory / FILE
    book.save(path)
    return path


def rendered(report: ValidationReport) -> list[str]:
    return [issue.render() for issue in report.issues]


def test_contract_is_the_briefs() -> None:
    assert COLUMNS == BRIEF_COLUMNS
    assert LEVELS == ("price", "yield_pct", "spread_bps", "bid", "ask")
    assert LETTERS == tuple("ABCDEFGHIJKLM")
    assert MEDIA_TYPE == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.mark.parametrize(
    "address",
    ["G2\n", "G" + "1" * 4301, "XFE2", "g2", "G02"],
    ids=["newline", "huge-row", "column-limit", "lowercase", "leading-zero"],
)
def test_shared_cell_parser_refuses_noncanonical_rv_coordinates(
    tmp_path: Path, address: str
) -> None:
    raw = write(tmp_path, [VALID]).read_bytes()
    stream = io.BytesIO()
    with ZipFile(io.BytesIO(raw)) as original, ZipFile(stream, "w") as rewritten:
        for member in original.infolist():
            content = original.read(member)
            if member.filename == "xl/worksheets/sheet1.xml":
                extra = f'<row r="2"><c r={quoteattr(address)}><v>999</v></c></row>'
                content = content.replace(
                    b"</sheetData>",
                    (extra + "</sheetData>").encode(),
                )
            rewritten.writestr(member, content)
    result = validate_rv_bytes(FILE, stream.getvalue())
    assert not result.ok
    assert "canonical" in result.issues[0].problem


def test_shared_cell_parser_preserves_native_bounds_and_separates_metadata() -> None:
    first = Span(locator="sheet:rv_quotes!A1", text="provider")
    last = Span(locator="sheet:rv_quotes!XFD1048576", text="ignored data")
    named = Span(locator="sheet:A%20B%21!B2", text="distinct sheet")
    merged = Span(locator="sheet:rv_quotes!merged:N2:N3", text="", data_type="metadata")
    comment = Span(locator="part:xl%2Fcomments%2Fcomment1.xml!G2\n", text="comment")
    grids, issues = sheet_cells((first, last, named, merged, comment), FILE)
    assert issues == []
    assert grids == {"rv_quotes": {(1, 1): first, (16384, 1048576): last}, "A B!": {(2, 2): named}}
    assert grids["rv_quotes"][16384, 1048576] is last
    _, repeated = sheet_cells((first, first), FILE)
    assert len(repeated) == 1 and "repeats" in repeated[0].problem


def test_shared_cell_parser_refuses_sheet_uri_aliases() -> None:
    canonical = Span(locator="sheet:rv_quotes!G2", text="98.75")
    alias = Span(locator="sheet:%72v_quotes!G2", text="999")
    _, issues = sheet_cells((canonical, alias), FILE)
    assert len(issues) == 1 and "canonical sheet URI" in issues[0].problem


def test_issue_renders_file_sheet_cell_problem_and_fix() -> None:
    issue = Issue(file="f.xlsx", sheet="rv_quotes", cell="B2", problem="p", fix="x")
    assert issue.render() == "f.xlsx:rv_quotes!B2: p — fix: x"


def test_two_valid_rows_are_ok(tmp_path: Path) -> None:
    second = VALID | {"quote_kind": "indicative_talk", "price": None, "bid": 97, "ask": 98.5}
    report = validate_rv_quotes(write(tmp_path, [VALID, second]))
    assert rendered(report) == []
    assert report.ok
    assert report.rows == 2
    assert report.lines() == ("OK: 2 rows",)


AS_OF_FIX = "enter the quote date as text YYYY-MM-DD (format the column as Text so Excel keeps it)"
NUMBER_FIX = (
    "enter a plain number: digits, an optional sign and one decimal point "
    "(no %, units, spaces or thousands separators)"
)
QUOTE_KIND_FIX = "enter observed or indicative_talk; leave out rows without a quote"
CURRENCY_FIX = "enter the three-letter upper-case ISO 4217 code, e.g. USD"
FACILITY_FIX = "enter the facility's id using letters, digits, - and _ only, e.g. TLB-2030"
PROVIDER_FIX = "enter the dealer, broker or pricing service that supplied the quote"
LEVEL_FIX = "enter at least one of price, yield_pct, spread_bps, bid or ask"


@pytest.mark.parametrize(
    "change, cell, problem, fix",
    [
        ({"provider": None}, "A2", "provider is empty", PROVIDER_FIX),
        ({"provider": "  "}, "A2", "provider is empty", PROVIDER_FIX),
        (
            {"provider": " example holdings LTD  (synthetic) "},
            "A2",
            "provider 'example holdings LTD  (synthetic)' is the issuer; a quote the issuer "
            "supplies is not third-party evidence",
            "name the third party that supplied the quote, or leave the issuer's own levels "
            "out of this sheet",
        ),
        ({"as_of": None}, "B2", "as_of is empty", AS_OF_FIX),
        (
            {"as_of": "30/09/2026"},
            "B2",
            "as_of '30/09/2026' is not an ISO date YYYY-MM-DD",
            AS_OF_FIX,
        ),
        (
            {"as_of": "2026-02-30"},
            "B2",
            "as_of '2026-02-30' is not an ISO date YYYY-MM-DD",
            AS_OF_FIX,
        ),
        (
            {"as_of": 46295},
            "B2",
            "as_of '46295' is an Excel date serial, not an ISO date YYYY-MM-DD",
            AS_OF_FIX,
        ),
        ({"issuer": None}, "C2", "issuer is empty", "enter the issuer's name"),
        ({"facility_id": None}, "D2", "facility_id is empty", FACILITY_FIX),
        (
            {"facility_id": "TLB 2030"},
            "D2",
            "facility_id 'TLB 2030' has characters other than letters, digits, - and _",
            FACILITY_FIX,
        ),
        (
            {"instrument": None},
            "E2",
            "instrument is empty",
            "describe the instrument, e.g. Term Loan B due 2030",
        ),
        ({"quote_kind": None}, "F2", "quote_kind is empty", QUOTE_KIND_FIX),
        (
            {"quote_kind": "none"},
            "F2",
            "quote_kind 'none' is not observed or indicative_talk",
            QUOTE_KIND_FIX,
        ),
        (
            {"quote_kind": "Observed"},
            "F2",
            "quote_kind 'Observed' is not observed or indicative_talk",
            QUOTE_KIND_FIX,
        ),
        ({"price": "1,000.5"}, "G2", "price '1,000.5' is not a number", NUMBER_FIX),
        (
            {"price": None, "yield_pct": "7.5%"},
            "H2",
            "yield_pct '7.5%' is not a number",
            NUMBER_FIX,
        ),
        (
            {"price": None, "spread_bps": "NaN"},
            "I2",
            "spread_bps 'NaN' is not a number",
            NUMBER_FIX,
        ),
        ({"price": None, "bid": "1_000"}, "J2", "bid '1_000' is not a number", NUMBER_FIX),
        ({"price": None, "ask": "abc"}, "K2", "ask 'abc' is not a number", NUMBER_FIX),
        ({"price": None}, "G2:K2", "no price, yield_pct, spread_bps, bid or ask", LEVEL_FIX),
        (
            {"bid": "99.50", "ask": "99.25"},
            "J2:K2",
            "bid 99.50 exceeds ask 99.25",
            "check the quote: bid must not exceed ask",
        ),
        ({"currency": None}, "L2", "currency is empty", CURRENCY_FIX),
        (
            {"currency": "usd"},
            "L2",
            "currency 'usd' is not a three-letter ISO 4217 code",
            CURRENCY_FIX,
        ),
        (
            {"currency": "US$"},
            "L2",
            "currency 'US$' is not a three-letter ISO 4217 code",
            CURRENCY_FIX,
        ),
        (
            {"price": "=97+1"},
            "G2",
            "price holds a formula; prepare freezes stored values and never recalculates",
            "replace the formula with its value (Paste Special, Values)",
        ),
    ],
)
def test_each_column_rule_fails_with_its_exact_message(
    tmp_path: Path, change: dict[str, Any], cell: str, problem: str, fix: str
) -> None:
    report = validate_rv_quotes(write(tmp_path, [VALID | change]))
    assert rendered(report) == [f"{FILE}:rv_quotes!{cell}: {problem} — fix: {fix}"]
    assert not report.ok
    assert report.rows == 1
    assert report.lines()[-1] == "FAILED: 1 issues"


def test_note_and_bid_equal_ask_are_accepted(tmp_path: Path) -> None:
    row = VALID | {"note": None, "price": None, "bid": "99.25", "ask": "99.250"}
    assert rendered(validate_rv_quotes(write(tmp_path, [row]))) == []


def test_stored_lexemes_are_read_as_decimals_not_floats(tmp_path: Path) -> None:
    # 0.1 + 0.2 > 0.3 in binary floating point; as decimals bid 0.3 does not exceed ask 0.3.
    row = VALID | {"price": None, "bid": "0.30000000000000000001", "ask": "0.3000000000000000001"}
    assert rendered(validate_rv_quotes(write(tmp_path, [row]))) == []
    flipped = VALID | {
        "price": None,
        "bid": "0.3000000000000000001",
        "ask": "0.30000000000000000001",
    }
    assert rendered(validate_rv_quotes(write(tmp_path, [flipped]))) == [
        f"{FILE}:rv_quotes!J2:K2: bid 0.3000000000000000001 exceeds ask 0.30000000000000000001"
        " — fix: check the quote: bid must not exceed ask"
    ]


def test_duplicate_quote_key_is_one_issue_on_the_later_row(tmp_path: Path) -> None:
    other = VALID | {"facility_id": "SN-2032", "instrument": "Senior Notes due 2032"}
    report = validate_rv_quotes(write(tmp_path, [VALID, other, VALID | {"price": 99}]))
    assert rendered(report) == [
        f"{FILE}:rv_quotes!A4:F4: duplicates row 2 on facility_id, as_of, provider and quote_kind"
        " — fix: keep one row per facility_id, as_of, provider and quote_kind, or correct the"
        " repeated value"
    ]
    assert report.rows == 3


def test_a_value_with_a_line_break_still_renders_as_one_line(tmp_path: Path) -> None:
    report = validate_rv_quotes(write(tmp_path, [VALID | {"currency": "US\nD\t"}]))
    assert report.issues[0].problem == "currency 'US\nD' is not a three-letter ISO 4217 code"
    assert rendered(report) == [
        f"{FILE}:rv_quotes!L2: currency 'US\\nD' is not a three-letter ISO 4217 code"
        f" — fix: {CURRENCY_FIX}"
    ]


@pytest.mark.parametrize(
    "change, cell, column",
    [
        ({"currency": "X" * 32000}, "L2", "currency"),
        ({"quote_kind": "q" * 32000}, "F2", "quote_kind"),
        ({"price": None, "bid": "9" * 32000, "ask": "1"}, "J2:K2", "bid"),
    ],
)
def test_a_long_cell_value_is_echoed_to_80_characters(
    tmp_path: Path, change: dict[str, Any], cell: str, column: str
) -> None:
    (issue,) = validate_rv_quotes(write(tmp_path, [VALID | change])).issues
    long = change[column]
    assert issue.cell == cell
    assert long[:80] + "…" in issue.problem
    assert long[:81] not in issue.problem
    # The 80-character echo plus the fixed problem and fix wording stay well under 300.
    assert len(issue.problem) < 200
    assert len(issue.render()) < 300
    assert issue.render().startswith(f"{FILE}:rv_quotes!{cell}: ")


def test_a_long_header_value_is_echoed_to_80_characters(tmp_path: Path) -> None:
    (issue,) = validate_rv_quotes(
        write(tmp_path, [VALID], header=("h" * 32000, *COLUMNS[1:]))
    ).issues
    assert issue.problem == f"header is '{'h' * 80}…', expected 'provider'"


def test_an_80_character_value_is_echoed_whole(tmp_path: Path) -> None:
    (issue,) = validate_rv_quotes(write(tmp_path, [VALID | {"currency": "X" * 80}])).issues
    assert issue.problem == f"currency '{'X' * 80}' is not a three-letter ISO 4217 code"


def test_rows_with_an_invalid_key_cell_are_not_also_duplicates(tmp_path: Path) -> None:
    row = VALID | {"provider": None}
    report = validate_rv_quotes(write(tmp_path, [row, row]))
    assert [issue.cell for issue in report.issues] == ["A2", "A3"]
    assert {issue.problem for issue in report.issues} == {"provider is empty"}


def test_blank_rows_between_quotes_are_skipped(tmp_path: Path) -> None:
    other = VALID | {"quote_kind": "indicative_talk"}
    empty = dict.fromkeys(COLUMNS)
    report = validate_rv_quotes(write(tmp_path, [VALID, empty, other]))
    assert rendered(report) == []
    assert report.rows == 2


def test_header_out_of_order_fails_on_the_header_not_on_rows(tmp_path: Path) -> None:
    swapped = ("as_of", "provider", *COLUMNS[2:])
    report = validate_rv_quotes(write(tmp_path, [VALID], header=swapped))
    assert rendered(report) == [
        f"{FILE}:rv_quotes!A1: header is 'as_of', expected 'provider' — fix: {HEADER_FIX}"
    ]


def test_missing_header_column_is_one_issue(tmp_path: Path) -> None:
    report = validate_rv_quotes(write(tmp_path, [], header=COLUMNS[:-1]))
    assert rendered(report) == [
        f"{FILE}:rv_quotes!M1: header is empty, expected 'note' — fix: {HEADER_FIX}"
    ]


def test_extra_header_column_is_one_issue(tmp_path: Path) -> None:
    report = validate_rv_quotes(write(tmp_path, [VALID], header=(*COLUMNS, "comment")))
    assert rendered(report) == [
        f"{FILE}:rv_quotes!N1: header is 'comment', expected nothing after M1 — fix: {HEADER_FIX}"
    ]


def test_header_only_sheet_has_no_quote_rows(tmp_path: Path) -> None:
    report = validate_rv_quotes(write(tmp_path, []))
    assert rendered(report) == [
        f"{FILE}:rv_quotes!A2: no quote rows below the header"
        " — fix: add one row per quote, starting in row 2"
    ]
    assert report.rows == 0


def test_empty_sheet_is_one_issue(tmp_path: Path) -> None:
    report = validate_rv_quotes(write(tmp_path, [], header=()))
    assert rendered(report) == [
        f"{FILE}:rv_quotes!A1: sheet rv_quotes is empty — fix: copy the header row from"
        " docs/templates/rv-quotes.template.xlsx and add one row per quote"
    ]


def test_missing_sheet_names_the_sheets_found(tmp_path: Path) -> None:
    report = validate_rv_quotes(write(tmp_path, [VALID], sheet="RV Quotes"))
    assert rendered(report) == [
        f"{FILE}:rv_quotes!A1: no sheet named rv_quotes (found: RV Quotes) — fix: name the"
        " quote sheet rv_quotes, as in docs/templates/rv-quotes.template.xlsx"
    ]
    assert report.rows == 0


def test_a_file_that_is_not_a_workbook_is_one_issue(tmp_path: Path) -> None:
    path = tmp_path / "quotes.csv"
    path.write_text("provider,as_of\nSynthetic Dealer A,2026-09-30\n")
    report = validate_rv_quotes(path)
    assert rendered(report) == [
        "quotes.csv:rv_quotes!A1: prepare cannot read this file as an OOXML workbook:"
        " Preparation failed (BadZipFile): File is not a zip file"
        " — fix: save the file as an Excel Workbook (.xlsx) and validate again"
    ]
    assert report.lines() == (*rendered(report), "FAILED: 1 issues")


def zip_with_workbook_member(payload: bytes, *, deflated: bool, encrypted: bool) -> bytes:
    """A one-member ZIP whose `xl/workbook.xml` header is patched to claim deflate
    compression or ZipCrypto encryption that its bytes do not honour."""
    stream = io.BytesIO()
    with ZipFile(stream, "w", ZIP_STORED) as archive:
        archive.writestr("xl/workbook.xml", payload)
    raw = bytearray(stream.getvalue())
    central = raw.index(b"PK\x01\x02")
    if deflated:  # compression method: local header offset 8, central directory offset 10
        raw[8:10] = raw[central + 10 : central + 12] = ZIP_DEFLATED.to_bytes(2, "little")
    if encrypted:  # general-purpose flag bit 0: local offset 6, central offset 8
        raw[6] |= 1
        raw[central + 8] |= 1
    return bytes(raw)


DAMAGED_FIX = (
    "save the file again as an Excel Workbook (.xlsx) and validate again; prepare would also "
    "refuse this file"
)


@pytest.mark.parametrize(
    "raw, detail",
    [
        (
            zip_with_workbook_member(b"\xff" * 64, deflated=True, encrypted=False),
            "zlib.error: Error -3 while decompressing data: invalid block type",
        ),
        (
            zip_with_workbook_member(b"<workbook/>", deflated=False, encrypted=True),
            "RuntimeError: File 'xl/workbook.xml' is encrypted, password required for extraction",
        ),
    ],
    ids=["corrupt-deflate", "encrypted-member"],
)
def test_a_damaged_zip_member_is_one_issue_not_a_traceback(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
    raw: bytes,
    detail: str,
) -> None:
    path = tmp_path / FILE
    path.write_bytes(raw)
    expected = (
        f"{FILE}:rv_quotes!A1: prepare cannot read this file as an OOXML workbook: damaged ZIP"
        f" member ({detail}) — fix: {DAMAGED_FIX}"
    )
    assert rendered(validate_rv_quotes(path)) == [expected]
    assert run_cli(["rv", str(path)], monkeypatch) == 1
    assert capsys.readouterr().out.splitlines() == [expected, "FAILED: 1 issues"]


def test_a_zip_that_is_not_ooxml_is_one_issue(tmp_path: Path) -> None:
    path = tmp_path / FILE
    with ZipFile(path, "w") as archive:
        archive.writestr("notes.txt", "synthetic, not a workbook")
    assert rendered(validate_rv_quotes(path)) == [
        f"{FILE}:rv_quotes!A1: prepare cannot read this file as an OOXML workbook: Preparation"
        " failed (KeyError): \"There is no item named 'xl/workbook.xml' in the archive\""
        " — fix: save the file as an Excel Workbook (.xlsx) and validate again"
    ]


def test_an_unreadable_path_is_one_issue(tmp_path: Path) -> None:
    report = validate_rv_quotes(tmp_path / "absent.xlsx")
    assert rendered(report) == [
        "absent.xlsx:rv_quotes!A1: cannot read the file: No such file or directory"
        " — fix: check the path and that the file is readable"
    ]


def test_report_ok_must_agree_with_issues() -> None:
    issue = Issue(file="f", sheet="rv_quotes", cell="A1", problem="p", fix="x")
    with pytest.raises(ValueError, match="ok"):
        ValidationReport(ok=True, rows=0, issues=(issue,))
    with pytest.raises(ValueError, match="ok"):
        ValidationReport(ok=False, rows=0, issues=())


def run_cli(arguments: list[str], monkeypatch: pytest.MonkeyPatch) -> int:
    from deploy_vii.__main__ import main

    monkeypatch.setattr(sys, "argv", ["deploy-vii", "validate-input", *arguments])
    with pytest.raises(SystemExit) as exit_info:
        main()
    code = exit_info.value.code
    assert isinstance(code, int)
    return code


def test_cli_prints_ok_and_exits_zero(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    path = write(tmp_path, [VALID, VALID | {"quote_kind": "indicative_talk"}])
    assert run_cli(["rv", str(path)], monkeypatch) == 0
    assert capsys.readouterr().out == "OK: 2 rows\n"


def test_cli_prints_one_line_per_issue_and_exits_one(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    path = write(tmp_path, [VALID | {"currency": "usd", "as_of": None}])
    assert run_cli(["rv", str(path)], monkeypatch) == 1
    assert capsys.readouterr().out.splitlines() == [
        f"{FILE}:rv_quotes!B2: as_of is empty — fix: {AS_OF_FIX}",
        f"{FILE}:rv_quotes!L2: currency 'usd' is not a three-letter ISO 4217 code"
        f" — fix: {CURRENCY_FIX}",
        "FAILED: 2 issues",
    ]


def test_cli_json_prints_the_report(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    path = write(tmp_path, [VALID | {"currency": "usd"}])
    assert run_cli(["rv", str(path), "--json"], monkeypatch) == 1
    printed = ValidationReport.model_validate_json(capsys.readouterr().out)
    assert printed == validate_rv_quotes(path)
    assert not printed.ok


def test_cli_kind_comes_from_the_input_table(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    from deploy_vii.inputs import KINDS

    assert tuple(KINDS) == ("rv", "portfolio")  # portfolio: task PF-1
    assert KINDS["rv"] is validate_rv_quotes
    assert run_cli(["quotes", str(tmp_path / "x.xlsx")], monkeypatch) == 2
    assert "invalid choice: 'quotes'" in capsys.readouterr().err


def test_template_validates_ok_with_six_synthetic_rows() -> None:
    report = validate_rv_quotes(TEMPLATE)
    assert rendered(report) == []
    assert report.lines() == ("OK: 6 rows",)
    assert {row["quote_kind"] for row in SYNTHETIC_ROWS} == {"observed", "indicative_talk"}
    assert len({row["facility_id"] for row in SYNTHETIC_ROWS}) == 2
    assert all("synthetic" in row["issuer"].casefold() for row in SYNTHETIC_ROWS)


def test_builder_is_byte_stable_and_matches_the_committed_template(tmp_path: Path) -> None:
    build(tmp_path / "first.xlsx")
    build(tmp_path / "second.xlsx")
    first = (tmp_path / "first.xlsx").read_bytes()
    assert first == (tmp_path / "second.xlsx").read_bytes()
    assert first == workbook_bytes(SYNTHETIC_ROWS)
    assert first == TEMPLATE.read_bytes(), "run: uv run python scripts/build_templates.py"


def test_an_independent_reader_sees_the_template_sheets_header_and_readme() -> None:
    book = openpyxl.load_workbook(TEMPLATE, read_only=True)
    assert book.sheetnames == ["rv_quotes", "README"]
    quotes = list(book["rv_quotes"].iter_rows(values_only=True))
    assert quotes[0] == COLUMNS
    assert len(quotes) == 7
    readme = [row[0] for row in book["README"].iter_rows(values_only=True)]
    assert readme == list(README_LINES)
    assert readme[0] == "synthetic example — replace every row"


def test_readme_documents_every_column_and_fix() -> None:
    readme = (ROOT / "docs" / "templates" / "README.md").read_text(encoding="utf-8")
    assert all(f"`{column}`" in readme for column in COLUMNS)
    assert [key for key, fix in FIX.items() if fix not in readme] == []
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in readme
    assert "sheet:rv_quotes" in readme
    assert "third_party" in readme


def test_prepare_freezes_the_template_as_a_ready_source_with_rv_quotes_locators(
    tmp_path: Path,
) -> None:
    shutil.copy(TEMPLATE, tmp_path / "rv-quotes.xlsx")
    run = Run(
        run_id="rv-template",
        issuer_id="example-holdings",
        issuer_name="Example Holdings Ltd (synthetic)",
        as_of=date(2026, 10, 15),
        case_kind="synthetic",
    )
    request = PreparationRequest(
        run=run,
        sources=(
            SourceRequest(source_id="rv-quotes", path="rv-quotes.xlsx", media_type=MEDIA_TYPE),
        ),
    )
    pack = prepare(request, tmp_path, tmp_path / "pack")
    assert load_pack(tmp_path / "pack") == pack
    (source,) = pack.sources
    assert source.status == "ready"
    quotes = {s.locator: s for s in source.spans if s.locator.startswith("sheet:rv_quotes!")}
    assert quotes["sheet:rv_quotes!A1"].text == "provider"
    assert [quotes[f"sheet:rv_quotes!{letter}1"].text for letter in LETTERS] == list(COLUMNS)
    assert quotes["sheet:rv_quotes!G2"].text == "98.75"
    assert quotes["sheet:rv_quotes!B2"].text == "2026-09-30"
    stored = sum(1 for row in SYNTHETIC_ROWS for column in COLUMNS if row.get(column))
    assert len(quotes) == len(COLUMNS) + stored
    assert not any(span.hidden or span.formula for span in quotes.values())


def test_a_file_too_large_to_prepare_is_one_issue(monkeypatch: pytest.MonkeyPatch) -> None:
    from deploy_vii import preparation

    monkeypatch.setattr(preparation, "MAX_BYTES", 4)
    assert rendered(validate_rv_quotes(TEMPLATE)) == [
        "rv-quotes.template.xlsx:rv_quotes!A1: prepare cannot read this file as an OOXML"
        " workbook: Preparation failed (ValueError): Source exceeds extraction byte limit"
        f" — fix: {FIX['workbook']}"
    ]


WORD = st.text(
    alphabet="abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789 .&()'-",
    min_size=1,
    max_size=24,
).filter(str.strip)
LEXEME = st.decimals(
    min_value=Decimal("-1000000"),
    max_value=Decimal("1000000"),
    places=4,
    allow_nan=False,
    allow_infinity=False,
).map(str)


@st.composite
def quote_rows(draw: st.DrawFn) -> list[dict[str, str]]:
    """Valid synthetic rows: distinct quote keys (the facility id carries the row index), a
    provider that is never the issuer, at least one numeric level and bid not above ask."""
    issuer = "Issuer " + draw(WORD)
    rows = []
    for index in range(draw(st.integers(min_value=1, max_value=6))):
        row = {
            "provider": "Provider " + draw(WORD),
            "as_of": draw(st.dates(date(2000, 1, 1), date(2099, 12, 31))).isoformat(),
            "issuer": issuer,
            "facility_id": f"F{index}-"
            + draw(st.from_regex(r"[A-Za-z0-9_-]{0,8}", fullmatch=True)),
            "instrument": draw(WORD),
            "quote_kind": draw(st.sampled_from(("observed", "indicative_talk"))),
            "currency": draw(st.from_regex(r"[A-Z]{3}", fullmatch=True)),
            "note": draw(st.one_of(st.just(""), WORD)),
        }
        for level in draw(st.lists(st.sampled_from(LEVELS), min_size=1, unique=True)):
            row[level] = draw(LEXEME)
        if "bid" in row and "ask" in row and Decimal(row["bid"]) > Decimal(row["ask"]):
            row["bid"], row["ask"] = row["ask"], row["bid"]
        rows.append(row)
    return rows


def openpyxl_bytes(rows: list[dict[str, str]]) -> bytes:
    book = openpyxl.Workbook()
    worksheet = book.active
    assert worksheet is not None
    worksheet.title = "rv_quotes"
    worksheet.append(list(COLUMNS))
    for row in rows:
        worksheet.append([row.get(column) or None for column in COLUMNS])
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


@settings(max_examples=60, deadline=None)
@given(quote_rows())
def test_valid_rows_written_through_the_builder_report_ok(rows: list[dict[str, str]]) -> None:
    report = validate_rv_bytes("p.xlsx", workbook_bytes(rows))
    assert rendered(report) == []
    assert report.rows == len(rows)
    # Differential: an independent writer (openpyxl, levels as text cells) reads the same.
    assert validate_rv_bytes("p.xlsx", openpyxl_bytes(rows)) == report


@settings(max_examples=60, deadline=None)
@given(quote_rows(), st.data())
def test_emptying_one_required_cell_is_exactly_one_issue_naming_it(
    rows: list[dict[str, str]], data: st.DataObject
) -> None:
    index = data.draw(st.integers(min_value=0, max_value=len(rows) - 1))
    column = data.draw(st.sampled_from(REQUIRED_COLUMNS))
    rows[index] = rows[index] | {column: ""}
    report = validate_rv_bytes("p.xlsx", workbook_bytes(rows))
    cell = f"{LETTERS[COLUMNS.index(column)]}{index + 2}"
    assert [(issue.cell, issue.problem) for issue in report.issues] == [
        (cell, f"{column} is empty")
    ]
    assert report.rows == len(rows)
